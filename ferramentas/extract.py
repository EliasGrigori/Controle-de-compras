# Extrai todas as abas da planilha (somente leitura) -> dados.json + dados_extraidos.csv
import sys, re, json, csv, unicodedata, difflib
import openpyxl, openpyxl.reader.excel as rx
rx.apply_stylesheet = lambda *a, **k: None   # a planilha tem estilos inválidos; ignoramos os estilos
SRC = '/workspace/Planilha de Compras Domésticas (2).xlsx'
OUT = '/workspace/controle-compras/'
wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)

def strip_acc(s): return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
MESES = ['JANEIRO','FEVEREIRO','MARCO','ABRIL','MAIO','JUNHO','JULHO','AGOSTO','SETEMBRO','OUTUBRO','NOVEMBRO','DEZEMBRO']
TYPO = {'JUNLO': 'JULHO', 'MAIL': 'MAIO', 'JANEIROS': 'JANEIRO'}

def parse_name(n):
    u = strip_acc(n).upper()
    day = int(re.match(r'\s*(\d{1,2})', u).group(1))
    ym = re.search(r'(20\d\d)', u); year = int(ym.group(1)) if ym else None
    month = None; how = 'exato'
    for w in re.findall(r'[A-Z]+', u):
        if w in MESES: month = MESES.index(w) + 1; break
        if w in TYPO: month = MESES.index(TYPO[w]) + 1; how = f'erro de digitação "{w}"→{TYPO[w].title()}'; break
    if month is None:
        for w in re.findall(r'[A-Z]{4,}', u):
            m = difflib.get_close_matches(w, MESES, 1, 0.75)
            if m: month = MESES.index(m[0]) + 1; how = f'aproximado "{w}"→{m[0]}'; break
    buyers = [b for b in ('Elias', 'Claudia', 'Juliano') if b.upper() in u]
    return day, month, year, how, ' + '.join(buyers)

# ---------- nomes canônicos (aliases) ----------
ALIAS = {
 'adocante gota':'Adoçante gotas','adocante po':'Adoçante em pó','adocante':'Adoçante em pó',
 'alho kg':'Alho','alho':'Alho','alho lata':'Alho (lata)',
 'amaciente':'Amaciante','acucar':'Açúcar','acafrao':'Açafrão',
 'bayggon':'Baygon (inseticida)','c. de leite':'Creme de leite','cacal em po':'Cacau em pó',
 'cafe':'Café','carbe mista 2,5kg':'Carne mista 2,5kg','carne de lata':'Carne de lata','carne enlatada':'Carne de lata',
 'carne hanbur.':'Carne de hambúrguer','carne moida':'Carne moída','carne muida':'Carne moída',
 'catchup':'Ketchup','cebola (kg)':'Cebola','cenosa(kg)':'Cenoura','continente':'Cotonete','contonete':'Cotonete',
 'coxa/sobrec':'Coxa e sobrecoxa','creme barbear':'Creme de barbear','cucuz':'Flocão de milho (cuscuz)',
 'ervilha':'Ervilha','ervilha/milho':'Ervilha/milho','m e ervilha':'Ervilha/milho',
 'essencia':'Essência (aromatizador)','essencia balnilha':'Essência de baunilha',
 'f. teigo c/ f.':'Farinha de trigo c/ fermento','f.trigo c/ ferm':'Farinha de trigo c/ fermento','farinha de trigo c/ferm':'Farinha de trigo c/ fermento',
 'f.trigo':'Farinha de trigo','farinha de trigo':'Farinha de trigo','farinha':'Farinha (mandioca)',
 'farinha lactea':'Farinha láctea','feijao':'Feijão','figado':'Fígado bovino','figado b. (kg)':'Fígado bovino','figado frango':'Fígado de frango',
 'formitol - veneno':'Formitol (veneno formiga)','formitoo':'Formitol (veneno formiga)','frango(kg)':'Frango','frango a passarinho':'Frango a passarinho',
 'gilete':'Gilete','iogurte':'Iogurte','leite cond.':'Leite condensado','leite condensado':'Leite condensado','leite em po':'Leite em pó',
 'lemon piper':'Lemon pepper','limon pepperr':'Lemon pepper','m. de tomate':'Molho de tomate',
 'mortadela':'Mortadela','miojo':'Miojo','nissin miojo':'Miojo','oleo':'Óleo de cozinha','óleo':'Óleo de cozinha','oleo mobile':'Óleo Mobil',
 'ovos(20un)':'Ovos (20 un)','ovos(30un)':'Ovos (30 un)','p. de dente m':'Pasta de dente M','p. de dente p':'Pasta de dente P','pasta de dente':'Pasta de dente M',
 'p. sanitaria':'Pedra sanitária','pedra sanitaria':'Pedra sanitária','pano micro fibra':'Pano de microfibra','papel aluminio':'Papel alumínio',
 'pegador se roupa':'Prendedor de roupa','phimichurri':'Chimichurri','pimentao':'Pimentão','pimentao(un)':'Pimentão','pao amburger':'Pão de hambúrguer',
 'recheado 4un':'Biscoito recheado','recheado':'Biscoito recheado','sabobete':'Sabonete','sabao liquido':'Sabão líquido','sabao em po ype':'Sabão em pó Ypê',
 'saso de lixo':'Saco de lixo','sal':'Sal','sentinela':'Sentinela','shampoo':'Shampoo','tenis pe baruel':'Tênis Pé Baruel (talco)',
 'temp. reg. feango':'Tempero regional frango','temp. folver. caseiro':'Tempero folhas verdes caseiro','temp edu':'Tempero Edu','temp gurme':'Tempero Gourmet',
 'temp. baiano':'Tempero baiano','vassoura':'Vassoura','veveno mosca':'Veneno para mosca','waffer':'Wafer','agua sanitaria':'Água sanitária',
 'hidratante':'Hidratante','h2oh limao':'H2OH limão','mistos de carnes e verduras':'Mistos de carnes e verduras','sabores':'Sabores (para dindin)',
 'sopao':'Sopão','salgado':'Salgado','molho shoyu':'Molho shoyu','batata palha':'Batata palha','batata doce':'Batata-doce',
}
def key(s): return strip_acc(re.sub(r'\s+', ' ', s.strip())).lower()
def nice(s):
    s = re.sub(r'\s+', ' ', s.strip())
    s = s.lower()
    return s[0].upper() + s[1:]
def canon(raw):
    k = key(raw)
    return ALIAS.get(k) or nice(raw)

# ---------- categorias ----------
CAT = {}
def setc(cat, names):
    for n in names: CAT[n] = cat
setc('Carne', ['Calabresa','Capa de filé','Carne alcatra','Carne de hambúrguer','Carne de lata','Carne mista 2,5kg','Carne moída','Costela bovina',
  'Coxa e sobrecoxa','Empanado','Fígado bovino','Fígado de frango','Frango','Frango a passarinho','Linguiça','Mortadela','Mortadela 1,5kg','Mortadela 1kg',
  'P. de frango','Peito de frango','Presunto','Sardinha'])
setc('Grãos', ['Amendoim','Arroz','Aveia','Ervilha','Ervilha/milho','Farinha (mandioca)','Farinha de trigo','Farinha de trigo c/ fermento','Feijão',
  'Flocão de milho (cuscuz)','Goma','Granola','Macarrão','Macarrão lasanha','Macarrão parafuso','Milho verde','Miojo','Pipoca','Sucrilhos'])
setc('Limpeza', ['Amaciante','Baygon (inseticida)','Bombril','Desinfetante','Detergente','Esponja','Essência (aromatizador)','Flanela','Formitol (veneno formiga)',
  'Naftalina','Pano de microfibra','Pasta mata barata','Pedra sanitária','Polidor','Refil difusor essência','Rodo','Sabão em barra','Sabão em pó',
  'Sabão em pó Ypê','Sabão líquido','Sabão pasta','Saco de lixo','Vassoura','Veja','Veneno para mosca','Água sanitária'])
UNIT = {'Alho':'kg','Cebola':'kg','Cenoura':'kg','Frango':'kg','Fígado bovino':'kg','Batata':'kg','Batata-doce':'kg','Carne moída':'kg','Tomate':'kg',
  'Costela bovina':'kg','Carne alcatra':'kg','Peito de frango':'kg','Coxa e sobrecoxa':'kg','Fígado de frango':'kg','Arroz':'kg','Feijão':'kg','Açúcar':'kg',
  'Farinha de trigo':'kg','Farinha de trigo c/ fermento':'kg','Óleo de cozinha':'un','Água sanitária':'L','Desinfetante':'L','Leite':'L','Amaciante':'un',
  'Biscoito recheado':'pct','Macarrão':'pct','Miojo':'pct','Frisco':'pct','Flocão de milho (cuscuz)':'pct','Goma':'kg','Ovos (20 un)':'cart','Ovos (30 un)':'cart','Saco de lixo':'pct'}

CATK = {key(k): v for k, v in CAT.items()}
BOUGHT_MARKS = {'👌','x','-','a','&','♥️','🔒','y','feito','g','x30','não','nao'}
def num(v):
    if isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    return None

compras = []; lines = []; issues = []
cur_y = cur_m = None
for idx, n in enumerate(wb.sheetnames):
    rows = list(wb[n].iter_rows(values_only=True))
    if not rows:
        issues.append(f'Aba "{n}" vazia — ignorada'); continue
    day, month, year, how, buyer = parse_name(n)
    inferred = False
    if year is None:
        year = cur_y if month <= cur_m else cur_y - 1; inferred = True
    cur_y, cur_m = year, month
    date = f'{year:04d}-{month:02d}-{day:02d}'
    hdr_total = num(rows[1][3]) if len(rows) > 1 else None
    pnote = rows[1][4] if len(rows) > 1 and len(rows[1]) > 4 and rows[1][4] else ''
    obs = []
    if 'FEIRA' in n.upper(): obs.append('Feira de junho (comprada em 31/05)')
    if pnote: obs.append(str(pnote).strip())
    cid = f'c{date.replace("-","")}_{idx:02d}'
    c = dict(id=cid, data=date, comprador=buyer, fornecedor='', obs='; '.join(obs), origem=n.strip(), anoInferido=inferred,
             mesInterpretado=how, totalPlanilha=round(hdr_total, 2) if hdr_total is not None else None, linhas=[])
    if inferred: issues.append(f'"{n.strip()}": ano não consta no nome → inferido {year} pela ordem das abas')
    if how != 'exato': issues.append(f'"{n.strip()}": mês com {how}')
    if not buyer: issues.append(f'"{n.strip()}": sem nome de comprador na aba → comprador vazio')
    if hdr_total == 0: issues.append(f'"{n.strip()}": total da planilha = 0 (lista de compras sem preços anotados)')
    for r in rows[3:]:
        r = list(r) + [None] * 6
        name, q, p, t, note = r[0], r[1], r[2], r[3], r[4]
        if not isinstance(name, str) or not name.strip(): continue
        qn, pn, tn = num(q), num(p), num(t)
        marks = []
        for v in (q, p, t):
            if isinstance(v, str) and v.strip(): marks.append(v.strip())
        qty = qn
        if isinstance(q, str) and re.fullmatch(r'\d+\s*g', q.strip().lower()):
            qty = float(re.match(r'\d+', q.strip()).group()) / 1000; issues.append(f'"{n.strip()}" / {name.strip()}: quantidade "{q}" convertida para {qty} kg')
        if isinstance(p, str) and re.fullmatch(r'\d+[.,]\d+', p.strip()):
            pn = float(p.strip().replace(',', '.')); issues.append(f'"{n.strip()}" / {name.strip()}: preço unitário digitado como texto "{p}" → {pn}; total não somado (planilha marca 👌)')
        lowmarks = {m.lower() for m in marks}
        bought_mark = bool(lowmarks & BOUGHT_MARKS) or any(m in ('?', '!') for m in marks if isinstance(q, str))
        has = (qty or 0) > 0 or (pn or 0) > 0 or (tn or 0) > 0 or bought_mark
        if not has:
            if marks: issues.append(f'"{n.strip()}" / {name.strip()}: só marca "{"/".join(marks)}" sem quantidade/preço → ignorada')
            continue
        prod = canon(name)
        cat = CATK.get(key(prod), 'Outros')
        sem_preco = not ((tn or 0) > 0)
        mk = '/'.join(m for m in marks if m not in (q if isinstance(q, str) and re.fullmatch(r'\d+\s*g', q.strip().lower()) else '',))
        L = dict(produto=prod, original=re.sub(r'\s+', ' ', name.strip()), qtd=qty if qty else None, preco=pn if pn else None,
                 total=round(tn, 2) if tn else None, marca=mk, obs=str(note).strip() if note else '')
        c['linhas'].append(L)
        lines.append(dict(data=date, planilha=n.strip(), comprador=buyer, categoria=cat, **L))
    s = round(sum(l['total'] or 0 for l in c['linhas']), 2)
    c['totalCalc'] = s
    if hdr_total is not None and abs(s - hdr_total) > 0.01: issues.append(f'"{n.strip()}": soma {s} ≠ total da planilha {hdr_total}')
    compras.append(c)

# itens
prods = sorted({l['produto'] for l in lines}, key=lambda s: strip_acc(s).lower())
items = []
for i, p in enumerate(prods):
    items.append(dict(id=f'i{i+1:03d}', nome=p, categoria=CATK.get(key(p), 'Outros'), unidade={key(k):v for k,v in UNIT.items()}.get(key(p), 'un'), fornecedor='', obs=''))
pid = {it['nome']: it['id'] for it in items}
unit_of = {it['nome']: it['unidade'] for it in items}
for c in compras:
    for L in c['linhas']:
        if unit_of[L['produto']] == 'kg' and L['qtd'] and L['qtd'] >= 100 and not L['preco']:
            issues.append(f'"{c["origem"]}" / {L["original"]}: quantidade {L["qtd"]:g} interpretada como gramas → {L["qtd"]/1000:g} kg')
            L['qtd'] = L['qtd'] / 1000
for l in lines:
    if unit_of[l['produto']] == 'kg' and l['qtd'] and l['qtd'] >= 100 and not l['preco']: l['qtd'] = l['qtd'] / 1000
for c in compras:
    for L in c['linhas']: L['itemId'] = pid[L.pop('produto')]
unused = sorted(set(CATK) - {key(p) for p in prods})
data = dict(versao=1, geradoEm='2026-10-03', origem='Planilha de Compras Domésticas (2).xlsx', itens=items, compras=sorted(compras, key=lambda c: c['data'], reverse=True),
            fornecedores=[], compradores=['Elias', 'Claudia', 'Juliano'])
json.dump(data, open(OUT + '_build/dados.json', 'w'), ensure_ascii=False, separators=(',', ':'))
json.dump(dict(issues=issues, unused=unused), open(OUT + '_build/issues.json', 'w'), ensure_ascii=False, indent=1)

def br(v, d=2):
    if v is None: return ''
    s = f'{v:.{d}f}'.rstrip('0').rstrip('.') if d == 3 else f'{v:.2f}'
    return s.replace('.', ',')
with open(OUT + 'dados_extraidos.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f, delimiter=';')
    w.writerow(['data', 'planilha', 'comprador', 'produto', 'produto_original', 'categoria', 'quantidade', 'preco_unitario', 'preco_total', 'marcacao', 'observacao'])
    for l in sorted(lines, key=lambda l: (l['data'], strip_acc(l['produto']).lower())):
        w.writerow([l['data'], l['planilha'], l['comprador'], l['produto'], l['original'], l['categoria'], br(l['qtd'], 3), br(l['preco']), br(l['total']), l['marca'], l['obs']])

from collections import defaultdict
st = defaultdict(lambda: [0, 0.0, set()])
for l in lines:
    s = st[l['categoria']]; s[0] += 1; s[1] += l['total'] or 0; s[2].add(l['produto'])
print('abas:', len(compras), 'datas:', min(c['data'] for c in compras), '→', max(c['data'] for c in compras))
print('linhas:', len(lines), 'produtos:', len(prods), 'total R$', round(sum(l['total'] or 0 for l in lines), 2))
for k in ('Carne', 'Grãos', 'Limpeza', 'Outros'): print(k, st[k][0], 'linhas', round(st[k][1], 2), 'R$', len(st[k][2]), 'produtos')
print('categoria sem uso:', unused)
for c in compras: print(c['data'], c['comprador'] or '-', len(c['linhas']), c['totalCalc'], '|', c['origem'])
